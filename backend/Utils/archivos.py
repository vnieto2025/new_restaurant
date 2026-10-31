import csv
import io
from datetime import date, datetime
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from fastapi.responses import Response

TIPOS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
# Una hoja de cálculo ejecuta como fórmula el texto que empieza así: un cliente llamado "=HYPERLINK(...)" no debe poder hacerlo
INICIOS_DE_FORMULA = ("=", "+", "-", "@", "\t", "\r")
ANCHO_MAXIMO = 45


def texto_csv(valor):
    """ Valor como lo espera un Excel en español con separador ";": coma decimal y texto que no se ejecuta """
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, (int, Decimal, float)):
        return str(valor).replace(".", ",")
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor)
    return f"'{texto}" if texto.startswith(INICIOS_DE_FORMULA) else texto


def generar_csv(titulos: list, filas: list):
    salida = io.StringIO(newline="")
    escritor = csv.writer(salida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(titulos)
    for fila in filas:
        escritor.writerow([texto_csv(valor) for valor in fila])
    # La marca BOM hace que Excel reconozca las tildes
    return ("﻿" + salida.getvalue()).encode("utf-8")


def celda_excel(hoja, valor):
    if isinstance(valor, bool):
        valor = "Sí" if valor else "No"
    elif isinstance(valor, Decimal):
        valor = float(valor)
    celda = WriteOnlyCell(hoja)
    if isinstance(valor, str):
        # Tipo texto explícito: lo que empiece por "=" se guarda como texto, nunca como fórmula
        celda._value, celda.data_type = ILLEGAL_CHARACTERS_RE.sub("", valor), "s"
        return celda
    celda.value = valor
    if isinstance(valor, datetime):
        celda.number_format = "yyyy-mm-dd hh:mm"
    elif isinstance(valor, date):
        celda.number_format = "yyyy-mm-dd"
    return celda


def generar_xlsx(titulo_hoja: str, titulos: list, filas: list):
    libro = Workbook(write_only=True)
    hoja = libro.create_sheet(titulo_hoja[:31])
    hoja.freeze_panes = "A2"
    anchos = [len(titulo) for titulo in titulos]
    for fila in filas[:200]:
        for posicion, valor in enumerate(fila):
            anchos[posicion] = max(anchos[posicion], len(texto_csv(valor)))
    for posicion, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(posicion)].width = min(ancho + 2, ANCHO_MAXIMO)
    encabezado = []
    for titulo in titulos:
        celda = celda_excel(hoja, titulo)
        celda.font = Font(bold=True)
        encabezado.append(celda)
    hoja.append(encabezado)
    for fila in filas:
        hoja.append([celda_excel(hoja, valor) for valor in fila])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def respuesta_archivo(nombre: str, formato: str, contenido: bytes):
    return Response(content=contenido, media_type=TIPOS[formato],
                    headers={"Content-Disposition": f'attachment; filename="{nombre}.{formato}"'})
