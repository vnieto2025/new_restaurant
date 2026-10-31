"""
Exportación de datos (docs 26, 26B y 41): CSV y Excel generados en el servidor, con permisos por módulo y por datos
sensibles, auditoría de cada descarga, texto que no se ejecuta como fórmula y sin copiar el archivo al registro.
"""
import io
from datetime import date, timedelta
from openpyxl import load_workbook
from conftest import key, product_id, table, pay_table
from Utils.archivos import generar_csv, generar_xlsx
from Config.db import SessionLocal
from Models.log_peticion_model import LogPeticionModel

# Hasta mañana: la jornada del restaurante puede ir en otro día que el reloj de quien corre la prueba
PERIODO = {"desde": str(date.today() - timedelta(days=5)), "hasta": str(date.today() + timedelta(days=1))}


def bajar(api, conjunto, formato="csv", **extra):
    return api.post("/exportacion/generar", {"conjunto": conjunto, "formato": formato, **extra})


def lineas(respuesta):
    return respuesta.content.decode("utf-8-sig").splitlines()


def test_formulas_no_se_ejecutan_y_formatos():
    titulos = ["A", "B", "C", "D", "E"]
    filas = [['=HYPERLINK("http://x")', -5, 1.5, "+573001234567", "con\x01control"]]
    csv = generar_csv(titulos, filas).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv and ";-5;1,5;" in csv and "'+57300" in csv

    hoja = load_workbook(io.BytesIO(generar_xlsx("Prueba", titulos, filas))).active
    assert hoja["A2"].data_type == "s" and hoja["A2"].value.startswith("=HYPERLINK")
    assert hoja["B2"].value == -5 and hoja["C2"].value == 1.5 and hoja["E2"].value == "concontrol"


def test_ventas_completas_y_auditadas(esquina):
    admin, _ = esquina["admin.esquina"]
    caja, _ = esquina["caja.esquina"]
    mesero, _ = esquina["mesero.salon"]
    mesa = table(mesero, "Mesa 4")
    mesero.data("/orden/create", {"mesa_id": mesa["id"], "clave_idempotencia": key(),
                                  "items": [{"producto_id": product_id(admin, "Gaseosa 400 ml"), "cantidad": 2}]}, expect=201)
    pay_table(caja, mesa_id=mesa["id"])

    claves = {c["clave"] for g in admin.data("/exportacion/catalogo")["grupos"] for c in g["conjuntos"]}
    assert {"ventas_productos", "cuentas", "pagos", "clientes", "horas", "liquidaciones", "entregas"} <= claves

    csv = bajar(admin, "ventas_productos", **PERIODO)
    assert csv.status_code == 200 and csv.headers["content-type"].startswith("text/csv")
    assert f"ventas_productos_{PERIODO['desde']}" in csv.headers["content-disposition"]
    assert lineas(csv)[0] == "Producto;Unidades;Valor" and any("Gaseosa 400 ml" in l for l in lineas(csv))

    hoja = load_workbook(io.BytesIO(bajar(admin, "cuentas", "xlsx", **PERIODO).content)).active
    assert hoja["A1"].value == "Jornada" and hoja.max_row >= 2 and isinstance(hoja["L2"].value, int)

    registro = admin.data("/auditoria/list", {"accion": "datos_exportados"})
    assert any(r["accion"] == "datos_exportados" for r in registro)


def test_todos_los_conjuntos_se_generan(esquina):
    admin, _ = esquina["admin.esquina"]
    for grupo in admin.data("/exportacion/catalogo")["grupos"]:
        for conjunto in grupo["conjuntos"]:
            for formato in ("csv", "xlsx"):
                respuesta = bajar(admin, conjunto["clave"], formato, **PERIODO)
                assert respuesta.status_code in (200, 409), (conjunto["clave"], respuesta.text)


def test_permisos_y_modulos(esquina, napoli):
    caja, _ = esquina["caja.esquina"]
    assert caja.post("/exportacion/catalogo").status_code == 403
    assert bajar(caja, "clientes").status_code == 403

    # Sin el módulo no se ofrece ni se entrega
    pizzeria, _ = napoli["napoli"]
    claves = {c["clave"] for g in pizzeria.data("/exportacion/catalogo")["grupos"] for c in g["conjuntos"]}
    assert "cuentas" not in claves and "existencias" not in claves and "ordenes" in claves
    assert bajar(pizzeria, "cuentas", **PERIODO).status_code == 403

    admin, _ = esquina["admin.esquina"]
    assert bajar(admin, "inventado").status_code == 403
    assert bajar(admin, "productos", "pdf").status_code == 400
    assert bajar(admin, "ventas_productos", desde="2020-02-01", hasta="2020-01-01").status_code == 400
    # Un periodo sin movimiento avisa en vez de bajar un archivo vacío
    vacio = bajar(admin, "ventas_productos", desde="2001-01-01", hasta="2001-01-02")
    assert vacio.status_code == 409 and "No hay datos" in vacio.json()["message"]


def test_aisla_empresas_y_no_copia_datos_al_registro(esquina, napoli):
    admin, _ = esquina["admin.esquina"]
    pizzeria, _ = napoli["napoli"]
    assert not any("Hamburguesa clásica" in l for l in lineas(bajar(pizzeria, "productos")))
    assert any("Hamburguesa clásica" in l for l in lineas(bajar(admin, "productos")))

    admin.post("/cliente/create", {"nombre": "Persona Privada", "telefono": "3175550101", "acepta_datos": True}, expect=201)
    assert any("Persona Privada" in l for l in lineas(bajar(admin, "clientes")))
    db = SessionLocal()
    try:
        ultimo = db.query(LogPeticionModel).filter(LogPeticionModel.service == "/exportacion/generar").order_by(LogPeticionModel.id.desc()).first()
        assert "clientes" in ultimo.response and "Persona Privada" not in ultimo.response
    finally:
        db.close()
