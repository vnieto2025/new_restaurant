"""
Importación de productos desde una hoja de cálculo (docs 12 y 41): plantilla, revisión sin guardar, errores por fila,
categorías nuevas, todo o nada al aplicar, actualización con auditoría de precio y permisos.
"""
import base64
import io
from openpyxl import Workbook, load_workbook
from conftest import find
from Class.ImportarProductos import leer_precio
from Config.db import SessionLocal
from Models.log_peticion_model import LogPeticionModel

TITULOS = ["Producto", "Categoría", "Precio", "Descripción"]


def xlsx(filas, titulos=TITULOS):
    libro = Workbook()
    libro.active.append(titulos)
    for fila in filas:
        libro.active.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return {"nombre": "carta.xlsx", "archivo": base64.b64encode(salida.getvalue()).decode()}


def csv(texto, codificacion="utf-8-sig"):
    return {"nombre": "carta.csv", "archivo": base64.b64encode(texto.encode(codificacion)).decode()}


def test_leer_precio():
    for entrada, esperado in [(18000, 18000), (18000.0, 18000), ("18000", 18000), ("18.000", 18000), ("18,000", 18000),
                              ("$ 1.250.000", 1250000), ("0", 0)]:
        assert leer_precio(entrada) == esperado, entrada
    for entrada in [None, "", "gratis", "18,5", 18.5, "18.00", True, "-5"]:
        assert leer_precio(entrada) is None, entrada


def test_plantilla_es_un_excel_que_se_puede_volver_a_subir(napoli):
    admin, _ = napoli["napoli"]
    plantilla = admin.post("/producto/importar/plantilla")
    assert plantilla.status_code == 200 and "plantilla_productos.xlsx" in plantilla.headers["content-disposition"]
    assert load_workbook(io.BytesIO(plantilla.content)).active["A1"].value == "Producto"
    revision = admin.data("/producto/importar/revisar", {"nombre": "p.xlsx", "archivo": base64.b64encode(plantilla.content).decode()})
    assert revision["total"] == 2


def test_revisar_no_guarda_y_aplicar_crea_todo(napoli):
    admin, _ = napoli["napoli"]
    archivo = xlsx([["Lasaña de prueba", "Pastas de prueba", 28000, "Con ragú"], ["Tiramisú de prueba", "Postres de prueba", "$ 12.500", None],
                    [None, None, None, None]])
    antes = len(admin.data("/producto/list", {"todos": True}))
    revision = admin.data("/producto/importar/revisar", archivo)
    assert revision["crear"] == 2 and revision["errores"] == 0 and revision["total"] == 2
    assert revision["categorias_nuevas"] == ["Pastas de prueba", "Postres de prueba"]
    assert len(admin.data("/producto/list", {"todos": True})) == antes

    resultado = admin.data("/producto/importar/aplicar", archivo)
    assert resultado["crear"] == 2
    productos = admin.data("/producto/list", {"todos": True})
    lasana = find(productos, nombre="Lasaña de prueba")
    assert lasana["precio"] == 28000 and lasana["descripcion"] == "Con ragú" and find(productos, nombre="Tiramisú de prueba")["precio"] == 12500
    assert find(admin.data("/categoria/list"), nombre="Pastas de prueba")
    # Queda en la carta, como un producto creado a mano
    assert lasana["cartas"]
    assert any(r["accion"] == "productos_importados" for r in admin.data("/auditoria/list", {"accion": "productos_importados"}))

    # Subirlo de nuevo no duplica: lo existente se omite y no hay nada que aplicar
    otra = admin.data("/producto/importar/revisar", archivo)
    assert otra["omitir"] == 2 and otra["crear"] == 0
    assert admin.post("/producto/importar/aplicar", archivo).status_code == 409


def test_errores_por_fila_y_todo_o_nada(napoli):
    admin, _ = napoli["napoli"]
    archivo = xlsx([["Buena de prueba", "Pastas de prueba", 20000], ["", "Pastas de prueba", 1000], ["Sin precio", "Pastas de prueba", "gratis"],
                    ["Buena de prueba", "Pastas de prueba", 20000], ["Sin categoría", "", 5000]], ["Producto", "Categoría", "Precio"])
    revision = admin.data("/producto/importar/revisar", archivo)
    assert revision["errores"] == 4 and revision["crear"] == 1
    assert revision["filas"][1]["fila"] == 3 and "nombre" in revision["filas"][1]["errores"][0]
    assert "precio" in revision["filas"][2]["errores"][0] and "anterior" in revision["filas"][3]["errores"][0]

    rechazo = admin.post("/producto/importar/aplicar", archivo)
    assert rechazo.status_code == 409 and "errores" in rechazo.json()["message"]
    assert not [p for p in admin.data("/producto/list", {"todos": True}) if p["nombre"] == "Buena de prueba"]

    # Punto de producción e impuesto deben existir
    malo = xlsx([["Con punto", "Pastas de prueba", 1000, "Parrilla inexistente"]], ["Producto", "Categoría", "Precio", "Punto de producción"])
    assert "Parrilla" in admin.data("/producto/importar/revisar", malo)["filas"][0]["errores"][0]


def test_archivos_invalidos(napoli):
    admin, _ = napoli["napoli"]
    assert admin.post("/producto/importar/revisar", {"nombre": "carta.pdf", "archivo": "aGVsbG8="}).status_code == 400
    assert admin.post("/producto/importar/revisar", {"nombre": "carta.xlsx", "archivo": "aGVsbG8="}).status_code == 400
    assert admin.post("/producto/importar/revisar", {"nombre": "carta.csv", "archivo": "esto no es base64!"}).status_code == 400
    faltan = admin.post("/producto/importar/revisar", csv("Nombre;Valor\nPan;100"))
    assert faltan.status_code == 400 and "Producto" in faltan.json()["message"]
    assert admin.post("/producto/importar/revisar", csv("Producto;Categoría;Precio\n")).status_code == 400
    grande = {"nombre": "carta.csv", "archivo": "A" * (3 * 1024 * 1024)}
    assert "2 MB" in admin.post("/producto/importar/revisar", grande).json()["message"]
    # Las filas vacías del final, que Excel guarda con formato, no cuentan
    vacias = "\n".join(";;" for _ in range(1500))
    assert admin.data("/producto/importar/revisar", csv(f"Producto;Categoría;Precio\nPan prueba;Cat;100\n{vacias}"))["total"] == 1
    filas = "\n".join(f"P{i};Cat;100" for i in range(1001))
    assert "1000" in admin.post("/producto/importar/revisar", csv(f"Producto;Categoría;Precio\n{filas}")).json()["message"]


def test_csv_de_excel_y_actualizar_precios(napoli):
    admin, _ = napoli["napoli"]
    mesero, _ = napoli["mesero.napoli"]
    # Excel en español guarda en Windows-1252 y con punto y coma; también se acepta la coma
    latino = csv("Producto;Categoría;Precio\nJugo de piña prueba;Jugos de prueba;7.000\n", "cp1252")
    assert admin.data("/producto/importar/aplicar", latino)["crear"] == 1
    producto = find(admin.data("/producto/list", {"todos": True}), nombre="Jugo de piña prueba")

    con_comas = csv('Producto,Categoría,Precio\nJugo de piña prueba,Jugos de prueba,"8.000"\n')
    assert admin.data("/producto/importar/revisar", con_comas)["omitir"] == 1
    revision = admin.data("/producto/importar/revisar", {**con_comas, "existentes": "actualizar"})
    assert revision["actualizar"] == 1 and revision["filas"][0]["precio_anterior"] == 7000
    admin.data("/producto/importar/aplicar", {**con_comas, "existentes": "actualizar"})
    assert find(admin.data("/producto/list", {"todos": True}), id=producto["id"])["precio"] == 8000
    # El cambio de precio queda auditado igual que a mano (RG-14)
    assert any(r["entidad_id"] == producto["id"] for r in admin.data("/auditoria/list", {"accion": "precio_cambiado"}))

    # Permisos, y el archivo no se copia al registro de peticiones
    assert mesero.post("/producto/importar/revisar", latino).status_code == 403
    db = SessionLocal()
    try:
        ultimo = db.query(LogPeticionModel).filter(LogPeticionModel.service == "/producto/importar/revisar",
                                                   LogPeticionModel.request.contains("carta.csv")).order_by(LogPeticionModel.id.desc()).first()
        assert con_comas["archivo"] not in ultimo.request
    finally:
        db.close()
