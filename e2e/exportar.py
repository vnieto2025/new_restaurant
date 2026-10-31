"""Recorrido de exportación: el administrador descarga ventas en Excel y CSV con datos reales, confirma antes de bajar
datos de personas, la descarga queda en auditoría y el cajero no ve la pantalla. También en celular."""
import csv
import uuid
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
from comun import APP, OUT, api, contexto, login, shot, abrir_navegador, terminar


admin = api("/login", {"usuario": "admin.esquina", "password": "demo1234"})["token"]
caja = api("/login", {"usuario": "caja.esquina", "password": "demo1234"})["token"]
if not api("/caja/estado", {}, caja)["turno"]:
    api("/caja/turno/abrir", {"fondo_inicial": 100000}, caja)

# Una venta cobrada, y un cliente con un nombre que una hoja de cálculo intentaría ejecutar
estado = api("/caja/estado", {}, caja)
efectivo = next(m["id"] for m in estado["medios"] if m["nombre"] == "Efectivo")
producto = next(p["id"] for p in api("/producto/list", {"todos": True}, admin) if p["nombre"] == "Limonada natural")
externa = next(m["id"] for m in api("/mesa/list", {"incluir_externa": True}, caja) if m["tipo"] == "externa")
orden = api("/orden/create", {"mesa_id": externa, "clave_idempotencia": uuid.uuid4().hex, "cliente_nombre": "Exportación",
                              "items": [{"producto_id": producto, "cantidad": 3}]}, caja)
cuenta = api("/cuenta/abrir", {"sesion_mesa_id": orden["sesion_mesa_id"]}, caja)
api("/cuenta/pagar", {"cuenta_id": cuenta["id"], "clave_idempotencia": uuid.uuid4().hex,
                      "pagos": [{"medio_pago_id": efectivo, "monto": cuenta["saldo"]}]}, caja)
api("/cliente/create", {"nombre": '=HYPERLINK("http://malo")', "telefono": "3" + str(uuid.uuid4().int)[:9], "acepta_datos": True}, admin)


def bajar(page, nombre):
    with page.expect_download() as descarga:
        page.get_by_role("button", name=nombre).click()
    ruta = OUT / descarga.value.suggested_filename
    descarga.value.save_as(ruta)
    return ruta


with sync_playwright() as p:
    browser = abrir_navegador(p)

    page = contexto(browser, 1366, 900)
    login(page, "admin.esquina")
    page.goto(f"{APP}/admin")
    page.get_by_role("link", name="Exportar datos").click()
    expect(page.get_by_role("heading", name="Exportar datos")).to_be_visible()
    expect(page.get_by_text("Horas trabajadas")).to_be_visible()
    shot(page, "x01-exportar")

    excel = bajar(page, "Descargar Ventas por producto en Excel")
    assert excel.name.startswith("ventas_productos_") and excel.suffix == ".xlsx"
    hoja = load_workbook(excel).active
    filas = {f[0]: f for f in hoja.iter_rows(values_only=True)}
    assert filas["Producto"] == ("Producto", "Unidades", "Valor") and filas["Limonada natural"][1] >= 3
    expect(page.get_by_text("archivo descargado")).to_be_visible()

    plano = bajar(page, "Descargar Pagos recibidos en CSV")
    with open(plano, encoding="utf-8-sig", newline="") as archivo:
        lineas = list(csv.reader(archivo, delimiter=";"))
    assert lineas[0][3] == "Medio" and any(l[3] == "Efectivo" for l in lineas[1:])

    # Datos de personas: se advierte antes y el nombre malicioso no queda como fórmula
    page.get_by_role("button", name="Descargar Clientes en Excel").click()
    dialogo = page.get_by_role("dialog")
    expect(dialogo).to_contain_text("datos de personas")
    shot(page, "x02-confirmar-personales")
    with page.expect_download() as descarga:
        dialogo.get_by_role("button", name="Descargar", exact=True).click()
    ruta = OUT / descarga.value.suggested_filename
    descarga.value.save_as(ruta)
    celdas = [f[0] for f in load_workbook(ruta).active.iter_rows(min_row=2)]
    malicioso = next(c for c in celdas if str(c.value).startswith("=HYPERLINK"))
    assert malicioso.data_type == "s"

    # Otro rango: sin fechas no se puede pedir; un periodo vacío avisa sin bajar nada
    page.get_by_role("button", name="Otro rango").click()
    expect(page.get_by_role("button", name="Descargar Ventas por producto en CSV")).to_be_disabled()
    page.get_by_label("Desde").fill("2001-01-01")
    page.get_by_label("Hasta").fill("2001-01-31")
    page.get_by_role("button", name="Descargar Ventas por producto en CSV").click()
    expect(page.get_by_text("No hay datos para descargar en ese periodo")).to_be_visible()
    shot(page, "x03-periodo-vacio")

    # Queda en auditoría
    page.goto(f"{APP}/admin/auditoria")
    expect(page.locator("td", has_text="Datos exportados").first).to_be_visible()
    shot(page, "x04-auditoria")
    page.context.close()

    # El cajero no tiene la opción
    page = contexto(browser, 1280, 860)
    login(page, "caja.esquina")
    page.goto(f"{APP}/admin/exportar")
    expect(page.get_by_role("heading", name="Exportar datos")).to_have_count(0)
    page.context.close()

    # Celular
    page = contexto(browser, 390, 844, mobile=True)
    login(page, "admin.esquina")
    page.goto(f"{APP}/admin/exportar")
    expect(page.get_by_text("Ventas por producto")).to_be_visible()
    shot(page, "x05-exportar-celular")

    browser.close()

terminar(esperados=("status of 409",))
