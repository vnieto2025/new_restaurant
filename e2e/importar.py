"""Recorrido de importación de productos: descargar plantilla, subir un archivo con errores y ver el detalle por fila,
corregirlo e importar; los productos quedan en la carta del cliente. También actualizar precios y verlo en celular."""
import uuid
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, expect
from comun import APP, OUT, api, contexto, login, shot, abrir_navegador, terminar

marca = uuid.uuid4().hex[:4]
nuevo = f"Sancocho {marca}"
token_qr = next(q for q in api("/qr/list", {}, api("/login", {"usuario": "admin.esquina", "password": "demo1234"})["token"])
                if q["mesa"] == "Terraza 1")["url"].rsplit("/", 1)[1]


def hoja(nombre, filas):
    libro = Workbook()
    libro.active.append(["Producto", "Categoría", "Precio", "Descripción"])
    for fila in filas:
        libro.active.append(fila)
    ruta = OUT / nombre
    libro.save(ruta)
    return ruta


con_errores = hoja("con_errores.xlsx", [[nuevo, f"Sopas {marca}", 24000, "Con arroz y aguacate"], ["", f"Sopas {marca}", 1000, None],
                                        [f"Ajiaco {marca}", f"Sopas {marca}", "gratis", None]])
correcto = hoja("correcto.xlsx", [[nuevo, f"Sopas {marca}", 24000, "Con arroz y aguacate"], [f"Ajiaco {marca}", f"Sopas {marca}", "$ 26.000", None]])
precios = hoja("precios.xlsx", [[nuevo, f"Sopas {marca}", 25000, "Con arroz y aguacate"]])

with sync_playwright() as p:
    browser = abrir_navegador(p)
    page = contexto(browser, 1366, 900)
    login(page, "admin.esquina")
    page.goto(f"{APP}/admin/productos")
    page.get_by_role("button", name="Importar").click()
    dialogo = page.get_by_role("dialog")
    expect(dialogo).to_contain_text("Descarga la plantilla")
    with page.expect_download() as descarga:
        dialogo.get_by_role("button", name="Descargar plantilla").click()
    descarga.value.save_as(OUT / "plantilla.xlsx")
    assert load_workbook(OUT / "plantilla.xlsx").active["C1"].value == "Precio"
    shot(page, "i01-importar")

    # Con errores: se ve cada fila y no deja importar
    dialogo.get_by_label("Archivo de productos").set_input_files(con_errores)
    expect(dialogo.locator(".resumen")).to_contain_text("con errores")
    expect(dialogo.get_by_text("El precio debe ser un valor en pesos")).to_be_visible()
    expect(dialogo.get_by_role("button", name="Importar", exact=True)).to_be_disabled()
    shot(page, "i02-errores")

    # Corregido
    dialogo.get_by_label("Archivo de productos").set_input_files(correcto)
    expect(dialogo.locator(".resumen")).to_contain_text("2 nuevos")
    expect(dialogo).to_contain_text(f"Sopas {marca}")
    shot(page, "i03-revision")
    dialogo.get_by_role("button", name="Importar 2 productos").click()
    expect(page.get_by_text("Importación lista")).to_be_visible()
    expect(page.get_by_text(nuevo)).to_be_visible()
    shot(page, "i04-importados")

    # Actualizar precios: muestra el precio anterior
    page.get_by_role("button", name="Importar").click()
    dialogo.get_by_label("Archivo de productos").set_input_files(precios)
    expect(dialogo.locator(".resumen")).to_contain_text("1 sin cambios")
    dialogo.get_by_text("Actualizarlo con los datos").click()
    expect(dialogo.locator(".resumen")).to_contain_text("1 por actualizar")
    expect(dialogo).to_contain_text("antes $ 24.000")
    shot(page, "i05-actualizar-precio")
    dialogo.get_by_role("button", name="Importar 1 producto").click()
    expect(page.get_by_text("Actualizados: 1")).to_be_visible()
    page.context.close()

    # Quedan en la carta del cliente
    cel = contexto(browser, 390, 844, mobile=True)
    cel.goto(f"{APP}/q/{token_qr}")
    expect(cel.get_by_text(nuevo)).to_be_visible()
    cel.get_by_text(nuevo).scroll_into_view_if_needed()
    shot(cel, "i06-en-la-carta")
    cel.context.close()

    # Celular
    page = contexto(browser, 390, 844, mobile=True)
    login(page, "admin.esquina")
    page.goto(f"{APP}/admin/productos")
    page.get_by_role("button", name="Importar").click()
    page.get_by_role("dialog").get_by_label("Archivo de productos").set_input_files(con_errores)
    expect(page.get_by_role("dialog").locator(".resumen")).to_contain_text("con errores")
    shot(page, "i07-celular")

    browser.close()

terminar()
